"""
app.py — SpareTrack Management System
Extended from Bike Spare Parts Inventory & Sales System.
All original endpoints preserved + 20 new features added.
"""

from flask import Flask, request, jsonify, send_from_directory, send_file, Response
from flask_cors import CORS
import sqlite3, os, base64, uuid, io, shutil
from datetime import datetime, date, timedelta

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                    Paragraph, Spacer, HRFlowable)
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT, TA_LEFT
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
# On Android (Chaquopy) the app's python files live in a read-only asset
# location, so a writable per-app data directory is passed in via this env
# var (set by MainActivity before the interpreter starts). Desktop/Electron
# usage is unaffected — DATA_DIR just falls back to BASE_DIR as before.
DATA_DIR     = os.environ.get('SPARETRACK_DATA_DIR') or BASE_DIR
DB_PATH      = os.path.join(DATA_DIR, 'database.db')
IMAGES_DIR   = os.path.join(DATA_DIR, 'static', 'images')
FRONTEND_DIR = os.path.abspath(os.path.join(BASE_DIR, '..', 'frontend'))
BACKUP_DIR   = os.path.join(DATA_DIR, 'backups')

os.makedirs(IMAGES_DIR, exist_ok=True)
os.makedirs(BACKUP_DIR,  exist_ok=True)

app = Flask(__name__, static_folder=os.path.join(BASE_DIR, 'static'))
CORS(app, resources={r"/*": {"origins": "*"}})

# ── Database ──────────────────────────────────────────────────────────────────

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn

def init_db():
    conn = get_db()
    conn.executescript("""
        CREATE TABLE IF NOT EXISTS companies (
            id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL UNIQUE, logo_path TEXT
        );
        CREATE TABLE IF NOT EXISTS models (
            id INTEGER PRIMARY KEY AUTOINCREMENT, company_id INTEGER NOT NULL,
            model_name TEXT NOT NULL, image_path TEXT,
            FOREIGN KEY (company_id) REFERENCES companies(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS items (
            id INTEGER PRIMARY KEY AUTOINCREMENT, model_id INTEGER NOT NULL,
            name TEXT NOT NULL, image_path TEXT,
            mrp REAL NOT NULL DEFAULT 0 CHECK(mrp>=0),
            selling_price REAL NOT NULL DEFAULT 0 CHECK(selling_price>=0),
            quantity INTEGER NOT NULL DEFAULT 0 CHECK(quantity>=0),
            FOREIGN KEY (model_id) REFERENCES models(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS transactions (
            id INTEGER PRIMARY KEY AUTOINCREMENT, item_id INTEGER NOT NULL,
            type TEXT NOT NULL CHECK(type IN ('add','sale')),
            quantity INTEGER NOT NULL CHECK(quantity>0), date_time TEXT NOT NULL,
            FOREIGN KEY (item_id) REFERENCES items(id)
        );
        CREATE TABLE IF NOT EXISTS customers (
            id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, mobile TEXT NOT NULL,
            alt_mobile TEXT, email TEXT, address TEXT, city TEXT, notes TEXT,
            reg_date TEXT NOT NULL DEFAULT (date('now'))
        );
        CREATE TABLE IF NOT EXISTS bills (
            id INTEGER PRIMARY KEY AUTOINCREMENT, bill_number TEXT NOT NULL,
            bill_date TEXT NOT NULL DEFAULT (date('now')), customer_id INTEGER NOT NULL,
            subtotal REAL NOT NULL DEFAULT 0, discount REAL NOT NULL DEFAULT 0,
            grand_total REAL NOT NULL DEFAULT 0, amount_paid REAL NOT NULL DEFAULT 0,
            pending REAL NOT NULL DEFAULT 0,
            status TEXT NOT NULL DEFAULT 'pending' CHECK(status IN ('paid','partial','pending')),
            notes TEXT, created_at TEXT NOT NULL DEFAULT (datetime('now')),
            FOREIGN KEY (customer_id) REFERENCES customers(id)
        );
        CREATE TABLE IF NOT EXISTS bill_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT, bill_id INTEGER NOT NULL,
            item_name TEXT NOT NULL, quantity INTEGER NOT NULL DEFAULT 1 CHECK(quantity>0),
            unit_price REAL NOT NULL DEFAULT 0 CHECK(unit_price>=0), total REAL NOT NULL DEFAULT 0,
            FOREIGN KEY (bill_id) REFERENCES bills(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS payments (
            id INTEGER PRIMARY KEY AUTOINCREMENT, bill_id INTEGER NOT NULL,
            amount REAL NOT NULL CHECK(amount>0),
            method TEXT NOT NULL DEFAULT 'cash' CHECK(method IN ('cash','upi','card','bank','other')),
            note TEXT, paid_at TEXT NOT NULL DEFAULT (datetime('now')),
            FOREIGN KEY (bill_id) REFERENCES bills(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT);
        INSERT OR IGNORE INTO settings (key, value) VALUES
            ('shop_name','SpareTrack'),('shop_address',''),('shop_mobile',''),
            ('shop_email',''),('shop_gst',''),('shop_logo',''),
            ('bill_prefix','BILL'),('bill_counter','1000');
        CREATE INDEX IF NOT EXISTS idx_cust_mobile ON customers(mobile);
        CREATE INDEX IF NOT EXISTS idx_bills_cust  ON bills(customer_id);
        CREATE INDEX IF NOT EXISTS idx_bills_date  ON bills(bill_date);
        CREATE INDEX IF NOT EXISTS idx_bills_status ON bills(status);
    """)
    conn.commit(); conn.close()

def save_image(b64, pfx='img'):
    if not b64: return None
    try:
        if ',' in b64:
            hdr, data = b64.split(',',1)
            mime = hdr.split(';')[0].split(':')[-1]
        else:
            data, mime = b64, 'image/png'
        ext = mime.split('/')[-1]
        if ext in ('jpeg','jpg'): ext='jpg'
        elif ext not in ('png','gif','webp','bmp'): ext='png'
        fn = f"{pfx}_{uuid.uuid4().hex[:10]}.{ext}"
        with open(os.path.join(IMAGES_DIR, fn),'wb') as f:
            f.write(base64.b64decode(data))
        return f"/static/images/{fn}"
    except Exception as e:
        app.logger.error(f"save_image: {e}"); return None

def ok(data=None, code=200):
    if data is None: return jsonify({'success':True}), code
    return jsonify(data), code

def err(msg, code=400):
    return jsonify({'error':msg}), code

def gs(conn, key, default=''):
    row = conn.execute('SELECT value FROM settings WHERE key=?',(key,)).fetchone()
    return row['value'] if row else default

def next_bill_no(conn, bill_date):

    row = conn.execute("""
        SELECT COALESCE(MAX(CAST(bill_number AS INTEGER)), 0) + 1
        FROM bills
        WHERE bill_date = ?
    """, (bill_date,)).fetchone()

    return str(row[0])

# def next_bill_no(conn):
#     prefix  = gs(conn,'bill_prefix','BILL')
#     counter = int(gs(conn,'bill_counter','1000'))
#     num = f"{prefix}-{counter:05d}"
#     conn.execute("UPDATE settings SET value=? WHERE key='bill_counter'",(str(counter+1),))
#     return num

def bill_status(grand, paid):
    if paid<=0: return 'pending'
    if paid>=grand: return 'paid'
    return 'partial'

# ── Frontend serving ──────────────────────────────────────────────────────────

@app.route('/')
def root(): return send_file(os.path.join(FRONTEND_DIR,'index.html'))

@app.route('/static/images/<path:fn>')
def serve_img(fn): return send_from_directory(IMAGES_DIR, fn)

@app.route('/<path:fn>')
def serve_fe(fn):
    fp = os.path.join(FRONTEND_DIR, fn)
    return send_file(fp) if os.path.isfile(fp) else (jsonify({'error':'not found'}),404)

# ══════════════════════════════════════════════════════════════════════════════
# ORIGINAL ENDPOINTS (100% preserved)
# ══════════════════════════════════════════════════════════════════════════════

@app.route('/companies')
def get_companies():
    c=get_db(); rows=c.execute('SELECT * FROM companies ORDER BY name').fetchall(); c.close()
    return ok([dict(r) for r in rows])

@app.route('/add_company',methods=['POST'])
def add_company():
    d=request.get_json(silent=True) or {}; name=(d.get('name') or '').strip()
    if not name: return err('Company name required')
    logo=save_image(d.get('logo'),'logo'); c=get_db()
    try:
        cur=c.execute('INSERT INTO companies (name,logo_path) VALUES (?,?)',(name,logo))
        c.commit(); row=c.execute('SELECT * FROM companies WHERE id=?',(cur.lastrowid,)).fetchone()
        return ok(dict(row),201)
    except sqlite3.IntegrityError: return err('Company already exists',409)
    finally: c.close()

@app.route('/update_company/<int:i>',methods=['PUT'])
def update_company(i):
    d=request.get_json(silent=True) or {}; name=(d.get('name') or '').strip()
    if not name: return err('Name required'); c=get_db()
    try:
        ex=c.execute('SELECT * FROM companies WHERE id=?',(i,)).fetchone()
        if not ex: return err('Not found',404)
        logo=save_image(d.get('logo'),'logo') if d.get('logo') else ex['logo_path']
        c.execute('UPDATE companies SET name=?,logo_path=? WHERE id=?',(name,logo,i)); c.commit(); return ok()
    except sqlite3.IntegrityError: return err('Name already exists',409)
    finally: c.close()

@app.route('/delete_company/<int:i>',methods=['DELETE'])
def delete_company(i):
    c=get_db()
    try:
        if not c.execute('SELECT id FROM companies WHERE id=?',(i,)).fetchone(): return err('Not found',404)
        c.execute('DELETE FROM companies WHERE id=?',(i,)); c.commit(); return ok()
    finally: c.close()

@app.route('/models')
def get_all_models():
    c=get_db()
    rows=c.execute("SELECT m.*,co.name AS company_name FROM models m JOIN companies co ON m.company_id=co.id ORDER BY co.name,m.model_name").fetchall()
    c.close(); return ok([dict(r) for r in rows])

@app.route('/models/<int:cid>')
def get_models(cid):
    c=get_db(); rows=c.execute('SELECT * FROM models WHERE company_id=? ORDER BY model_name',(cid,)).fetchall(); c.close()
    return ok([dict(r) for r in rows])

@app.route('/add_model',methods=['POST'])
def add_model():
    d=request.get_json(silent=True) or {}; cid=d.get('company_id'); name=(d.get('model_name') or '').strip()
    if not cid: return err('company_id required')
    if not name: return err('model_name required')
    img=save_image(d.get('image'),'model'); c=get_db()
    try:
        cur=c.execute('INSERT INTO models (company_id,model_name,image_path) VALUES (?,?,?)',(int(cid),name,img))
        c.commit(); row=c.execute('SELECT * FROM models WHERE id=?',(cur.lastrowid,)).fetchone()
        return ok(dict(row),201)
    finally: c.close()

@app.route('/update_model/<int:i>',methods=['PUT'])
def update_model(i):
    d=request.get_json(silent=True) or {}; name=(d.get('model_name') or '').strip()
    if not name: return err('model_name required'); c=get_db()
    try:
        ex=c.execute('SELECT * FROM models WHERE id=?',(i,)).fetchone()
        if not ex: return err('Not found',404)
        img=save_image(d.get('image'),'model') if d.get('image') else ex['image_path']
        cid=int(d.get('company_id') or ex['company_id'])
        c.execute('UPDATE models SET model_name=?,image_path=?,company_id=? WHERE id=?',(name,img,cid,i))
        c.commit(); return ok()
    finally: c.close()

@app.route('/delete_model/<int:i>',methods=['DELETE'])
def delete_model(i):
    c=get_db()
    try:
        if not c.execute('SELECT id FROM models WHERE id=?',(i,)).fetchone(): return err('Not found',404)
        c.execute('DELETE FROM models WHERE id=?',(i,)); c.commit(); return ok()
    finally: c.close()

@app.route('/items')
def get_all_items():
    c=get_db()
    rows=c.execute("SELECT i.*,m.model_name,co.name AS company_name FROM items i JOIN models m ON i.model_id=m.id JOIN companies co ON m.company_id=co.id ORDER BY co.name,m.model_name,i.name").fetchall()
    c.close(); return ok([dict(r) for r in rows])

@app.route('/items/<int:mid>')
def get_items(mid):
    c=get_db(); rows=c.execute('SELECT * FROM items WHERE model_id=? ORDER BY name',(mid,)).fetchall(); c.close()
    return ok([dict(r) for r in rows])

@app.route('/add_item',methods=['POST'])
def add_item():
    d=request.get_json(silent=True) or {}; mid=d.get('model_id'); name=(d.get('name') or '').strip()
    if not mid: return err('model_id required')
    if not name: return err('Part name required')
    try:
        mrp=float(d.get('mrp',0)); sp=float(d.get('selling_price',0)); qty=int(d.get('quantity',0))
        if mrp<0 or sp<0 or qty<0: raise ValueError
    except (ValueError,TypeError): return err('Invalid numeric values')
    img=save_image(d.get('image'),'item'); c=get_db()
    try:
        cur=c.execute('INSERT INTO items (model_id,name,image_path,mrp,selling_price,quantity) VALUES (?,?,?,?,?,?)',(int(mid),name,img,mrp,sp,qty))
        c.commit(); row=c.execute('SELECT * FROM items WHERE id=?',(cur.lastrowid,)).fetchone()
        return ok(dict(row),201)
    finally: c.close()

@app.route('/update_item/<int:i>',methods=['PUT'])
def update_item(i):
    d=request.get_json(silent=True) or {}; name=(d.get('name') or '').strip()
    if not name: return err('Name required')
    try:
        mrp=float(d['mrp']); sp=float(d['selling_price']); qty=int(d['quantity'])
        if mrp<0 or sp<0 or qty<0: raise ValueError
    except (ValueError,TypeError): return err('Invalid numeric values')
    c=get_db()
    try:
        ex=c.execute('SELECT * FROM items WHERE id=?',(i,)).fetchone()
        if not ex: return err('Not found',404)
        img=save_image(d.get('image'),'item') if d.get('image') else ex['image_path']
        mid=int(d.get('model_id') or ex['model_id'])
        c.execute('UPDATE items SET name=?,image_path=?,mrp=?,selling_price=?,quantity=?,model_id=? WHERE id=?',(name,img,mrp,sp,qty,mid,i))
        c.commit(); return ok()
    finally: c.close()

@app.route('/delete_item/<int:i>',methods=['DELETE'])
def delete_item(i):
    c=get_db()
    try:
        if not c.execute('SELECT id FROM items WHERE id=?',(i,)).fetchone(): return err('Not found',404)
        c.execute('DELETE FROM items WHERE id=?',(i,)); c.commit(); return ok()
    finally: c.close()

@app.route('/add_stock',methods=['POST'])
def add_stock():
    d=request.get_json(silent=True) or {}; item_id=d.get('item_id')
    if not item_id: return err('item_id required')
    try:
        qty=int(d.get('quantity')); assert qty>0
    except: return err('quantity must be a positive integer')
    c=get_db()
    try:
        item=c.execute('SELECT * FROM items WHERE id=?',(item_id,)).fetchone()
        if not item: return err('Item not found',404)
        new_qty=item['quantity']+qty
        c.execute('UPDATE items SET quantity=? WHERE id=?',(new_qty,item_id))
        c.execute('INSERT INTO transactions (item_id,type,quantity,date_time) VALUES (?,?,?,?)',(item_id,'add',qty,datetime.now().isoformat(timespec='seconds')))
        c.commit(); return ok({'success':True,'new_quantity':new_qty})
    finally: c.close()

@app.route('/sell_item',methods=['POST'])
def sell_item():
    d=request.get_json(silent=True) or {}; item_id=d.get('item_id')
    if not item_id: return err('item_id required')
    try:
        qty=int(d.get('quantity')); assert qty>0
    except: return err('quantity must be a positive integer')
    c=get_db()
    try:
        item=c.execute('SELECT * FROM items WHERE id=?',(item_id,)).fetchone()
        if not item: return err('Item not found',404)
        if item['quantity']<qty: return err(f'Insufficient stock. Available: {item["quantity"]}',400)
        new_qty=item['quantity']-qty
        c.execute('UPDATE items SET quantity=? WHERE id=?',(new_qty,item_id))
        c.execute('INSERT INTO transactions (item_id,type,quantity,date_time) VALUES (?,?,?,?)',(item_id,'sale',qty,datetime.now().isoformat(timespec='seconds')))
        c.commit(); return ok({'success':True,'new_quantity':new_qty,'revenue':round(item['selling_price']*qty,2)})
    finally: c.close()

@app.route('/transactions')
def get_transactions():
    fd=request.args.get('date','').strip(); limit=int(request.args.get('limit',500)); c=get_db()
    try:
        base="SELECT t.*,i.name AS item_name,i.selling_price,i.mrp,m.model_name,co.name AS company_name FROM transactions t JOIN items i ON t.item_id=i.id JOIN models m ON i.model_id=m.id JOIN companies co ON m.company_id=co.id"
        rows=(c.execute(base+" WHERE date(t.date_time)=? ORDER BY t.date_time DESC",(fd,)) if fd else c.execute(base+" ORDER BY t.date_time DESC LIMIT ?",(limit,))).fetchall()
        return ok([dict(r) for r in rows])
    finally: c.close()

# ══════════════════════════════════════════════════════════════════════════════
# ENHANCED DASHBOARD
# ══════════════════════════════════════════════════════════════════════════════

@app.route('/dashboard')
def get_dashboard():
    today=date.today().isoformat(); mstart=date.today().replace(day=1).isoformat(); c=get_db()
    try:
        def sc(sql,p=()):
            v=c.execute(sql,p).fetchone()[0]; return v if v is not None else 0
        sold_today=sc("SELECT COALESCE(SUM(quantity),0) FROM transactions WHERE type='sale' AND date(date_time)=?",(today,))
        added_today=sc("SELECT COALESCE(SUM(quantity),0) FROM transactions WHERE type='add' AND date(date_time)=?",(today,))
        parts_rev=sc("SELECT COALESCE(SUM(t.quantity*i.selling_price),0) FROM transactions t JOIN items i ON t.item_id=i.id WHERE t.type='sale' AND date(t.date_time)=?",(today,))
        top_items=c.execute("SELECT i.name,SUM(t.quantity) AS total_sold,i.selling_price,SUM(t.quantity*i.selling_price) AS revenue FROM transactions t JOIN items i ON t.item_id=i.id WHERE t.type='sale' GROUP BY t.item_id ORDER BY total_sold DESC LIMIT 5").fetchall()
        low_stock=c.execute("SELECT i.*,m.model_name,co.name AS company_name FROM items i JOIN models m ON i.model_id=m.id JOIN companies co ON m.company_id=co.id WHERE i.quantity<5 ORDER BY i.quantity").fetchall()
        weekly_rev=c.execute("SELECT date(t.date_time) AS day,COALESCE(SUM(t.quantity*i.selling_price),0) AS revenue,COALESCE(SUM(t.quantity),0) AS units FROM transactions t JOIN items i ON t.item_id=i.id WHERE t.type='sale' AND date(t.date_time)>=date('now','-6 days') GROUP BY date(t.date_time) ORDER BY day").fetchall()
        monthly_chart=c.execute("SELECT strftime('%Y-%m',bill_date) AS month,COALESCE(SUM(grand_total),0) AS revenue,COUNT(*) AS bill_count FROM bills WHERE bill_date>=date('now','-6 months','start of month') GROUP BY strftime('%Y-%m',bill_date) ORDER BY month").fetchall()
        return ok({
            'sold_today':int(sold_today),'added_today':int(added_today),'revenue_today':round(float(parts_rev),2),
            'top_items':[dict(r) for r in top_items],'low_stock':[dict(r) for r in low_stock],
            'low_stock_count':len(low_stock),
            'total_companies':sc('SELECT COUNT(*) FROM companies'),'total_models':sc('SELECT COUNT(*) FROM models'),
            'total_items':sc('SELECT COUNT(*) FROM items'),'total_stock':sc('SELECT COALESCE(SUM(quantity),0) FROM items'),
            'weekly_revenue':[dict(r) for r in weekly_rev],'date':today,
            'total_customers':sc('SELECT COUNT(*) FROM customers'),'total_bills':sc('SELECT COUNT(*) FROM bills'),
            'bills_revenue':round(float(sc('SELECT COALESCE(SUM(grand_total),0) FROM bills')),2),
            'bills_paid':round(float(sc('SELECT COALESCE(SUM(amount_paid),0) FROM bills')),2),
            'bills_pending':round(float(sc("SELECT COALESCE(SUM(pending),0) FROM bills WHERE status!='paid'")),2),
            'today_bills_rev':round(float(sc("SELECT COALESCE(SUM(grand_total),0) FROM bills WHERE bill_date=?",(today,))),2),
            'today_bills_count':int(sc("SELECT COUNT(*) FROM bills WHERE bill_date=?",(today,))),
            'month_revenue':round(float(sc("SELECT COALESCE(SUM(grand_total),0) FROM bills WHERE bill_date>=?",(mstart,))),2),
            'pending_count':int(sc("SELECT COUNT(*) FROM bills WHERE status!='paid'")),
            'monthly_chart':[dict(r) for r in monthly_chart],
        })
    finally: c.close()

# ══════════════════════════════════════════════════════════════════════════════
# CUSTOMERS
# ══════════════════════════════════════════════════════════════════════════════

@app.route('/customers')
def get_customers():
    q=request.args.get('q','').strip(); c=get_db()
    try:
        if q:
            like=f'%{q}%'
            rows=c.execute("SELECT * FROM customers WHERE name LIKE ? OR mobile LIKE ? OR city LIKE ? ORDER BY name",(like,like,like)).fetchall()
        else:
            rows=c.execute('SELECT * FROM customers ORDER BY name').fetchall()
        result=[]
        for r in rows:
            cust=dict(r)
            stats=c.execute("SELECT COUNT(*) as bill_count,COALESCE(SUM(grand_total),0) as total_amount,COALESCE(SUM(amount_paid),0) as total_paid,COALESCE(SUM(pending),0) as total_pending FROM bills WHERE customer_id=?",(r['id'],)).fetchone()
            cust.update(dict(stats)); result.append(cust)
        return ok(result)
    finally: c.close()

@app.route('/customers/<int:i>')
def get_customer(i):
    c=get_db()
    try:
        cust=c.execute('SELECT * FROM customers WHERE id=?',(i,)).fetchone()
        if not cust: return err('Not found',404)
        result=dict(cust)
        # Full history across every year, newest bill first.
        bills=c.execute("SELECT b.* FROM bills b WHERE b.customer_id=? ORDER BY b.bill_date DESC, b.id DESC",(i,)).fetchall()
        bill_list=[dict(b) for b in bills]
        if bill_list:
            ids=[b['id'] for b in bill_list]
            placeholders=','.join('?'*len(ids))
            item_rows=c.execute(f"SELECT bill_id,item_name,quantity,unit_price,total FROM bill_items WHERE bill_id IN ({placeholders}) ORDER BY id",ids).fetchall()
            items_by_bill={}
            for r in item_rows:
                items_by_bill.setdefault(r['bill_id'],[]).append(dict(r))
            for b in bill_list:
                b['items']=items_by_bill.get(b['id'],[])
        stats=c.execute("SELECT COALESCE(SUM(grand_total),0) AS total_amount,COALESCE(SUM(amount_paid),0) AS total_paid,COALESCE(SUM(pending),0) AS total_pending,COUNT(*) AS bill_count FROM bills WHERE customer_id=?",(i,)).fetchone()
        result['bills']=bill_list; result.update(dict(stats)); return ok(result)
    finally: c.close()

@app.route('/add_customer',methods=['POST'])
def add_customer():
    d=request.get_json(silent=True) or {}; name=(d.get('name') or '').strip(); mob=(d.get('mobile') or '').strip()
    if not name: return err('Customer name required')
    if not mob: return err('Mobile number required')
    c=get_db()
    try:
        cur=c.execute("INSERT INTO customers (name,mobile,alt_mobile,email,address,city,notes,reg_date) VALUES (?,?,?,?,?,?,?,?)",
            (name,mob,(d.get('alt_mobile') or '').strip(),(d.get('email') or '').strip(),(d.get('address') or '').strip(),(d.get('city') or '').strip(),(d.get('notes') or '').strip(),d.get('reg_date') or date.today().isoformat()))
        c.commit(); row=c.execute('SELECT * FROM customers WHERE id=?',(cur.lastrowid,)).fetchone()
        return ok(dict(row),201)
    finally: c.close()

@app.route('/update_customer/<int:i>',methods=['PUT'])
def update_customer(i):
    d=request.get_json(silent=True) or {}; name=(d.get('name') or '').strip(); mob=(d.get('mobile') or '').strip()
    if not name: return err('Name required')
    if not mob: return err('Mobile required')
    c=get_db()
    try:
        if not c.execute('SELECT id FROM customers WHERE id=?',(i,)).fetchone(): return err('Not found',404)
        c.execute("UPDATE customers SET name=?,mobile=?,alt_mobile=?,email=?,address=?,city=?,notes=? WHERE id=?",
            (name,mob,(d.get('alt_mobile') or '').strip(),(d.get('email') or '').strip(),(d.get('address') or '').strip(),(d.get('city') or '').strip(),(d.get('notes') or '').strip(),i))
        c.commit(); return ok()
    finally: c.close()
@app.route('/delete_customer/<int:i>', methods=['DELETE'])
def delete_customer(i):

    c = get_db()

    try:

        customer = c.execute(
            "SELECT id FROM customers WHERE id=?",
            (i,)
        ).fetchone()

        if not customer:
            return err('Customer not found', 404)

        c.execute(
            "DELETE FROM payments WHERE bill_id IN (SELECT id FROM bills WHERE customer_id=?)",
            (i,)
        )

        c.execute(
            "DELETE FROM bill_items WHERE bill_id IN (SELECT id FROM bills WHERE customer_id=?)",
            (i,)
        )

        c.execute(
            "DELETE FROM bills WHERE customer_id=?",
            (i,)
        )

        c.execute(
            "DELETE FROM customers WHERE id=?",
            (i,)
        )

        c.commit()

        return ok({
            "message": "Customer deleted successfully"
        })

    finally:
        c.close()

# ══════════════════════════════════════════════════════════════════════════════
# BILLS
# ══════════════════════════════════════════════════════════════════════════════

@app.route('/bills')
def get_bills():
    status=request.args.get('status','').strip(); q=request.args.get('q','').strip(); c=get_db()
    try:
        sql="SELECT b.*,cu.name AS customer_name,cu.mobile AS customer_mobile FROM bills b JOIN customers cu ON b.customer_id=cu.id"
        params=[]; wheres=[]
        if status: wheres.append("b.status=?"); params.append(status)
        if q:
            wheres.append("(cu.name LIKE ? OR cu.mobile LIKE ? OR b.bill_number LIKE ?)")
            like=f'%{q}%'; params+=[like,like,like]
        if wheres: sql+=" WHERE "+" AND ".join(wheres)
        sql+=" ORDER BY b.created_at DESC"
        rows=c.execute(sql,params).fetchall(); return ok([dict(r) for r in rows])
    finally: c.close()

@app.route('/bills/<int:i>')
def get_bill(i):
    c=get_db()
    try:
        bill=c.execute("SELECT b.*,cu.name AS customer_name,cu.mobile AS customer_mobile,cu.email AS customer_email,cu.address AS customer_address,cu.city AS customer_city FROM bills b JOIN customers cu ON b.customer_id=cu.id WHERE b.id=?",(i,)).fetchone()
        if not bill: return err('Not found',404)
        result=dict(bill)
        result['items']=[dict(x) for x in c.execute('SELECT * FROM bill_items WHERE bill_id=? ORDER BY id',(i,)).fetchall()]
        result['payments']=[dict(x) for x in c.execute('SELECT * FROM payments WHERE bill_id=? ORDER BY paid_at',(i,)).fetchall()]
        return ok(result)
    finally: c.close()
@app.route('/add_bill', methods=['POST'])
def add_bill():

    d = request.get_json(silent=True) or {}

    c = get_db()

    bill_date = d.get('bill_date') or date.today().isoformat()

    bill_no = next_bill_no(c, bill_date)

    cid = d.get('customer_id')
    customer_name = (d.get('customer_name') or '').strip()
    customer_mobile = (d.get('customer_mobile') or '').strip()

    items_data = d.get('items', [])

    if not cid and not customer_name:
        return err('Customer name required')

    if not items_data:
        return err('At least one item required')

    try:
        discount = float(d.get('discount', 0))
        paid = float(d.get('amount_paid', 0))

        if discount < 0 or paid < 0:
            raise ValueError

    except:
        return err('Invalid numeric values')

    

    try:

        customer = None

        if not cid:

            if not customer_name:
                return err('Customer name required')

            row = c.execute(
                """
                SELECT id
                FROM customers
                WHERE name=? AND mobile=?
                """,
                (
                    customer_name,
                    customer_mobile
                )
            ).fetchone()

            if row:

                cid = row["id"]

            else:

                cur_customer = c.execute(
                    """
                    INSERT INTO customers
                    (
                        name,
                        mobile,
                        city,
                        address
                    )
                    VALUES (?,?,?,?)
                    """,
                    (
                        customer_name,
                        customer_mobile,
                        '',
                        ''
                    )
                )

                cid = cur_customer.lastrowid

        customer = c.execute(
            "SELECT * FROM customers WHERE id=?",
            (cid,)
        ).fetchone()

        if not customer:
            return err('Customer not found', 404)

        notes = (d.get('notes') or '').strip()

        subtotal = 0.0
        parsed = []

        for it in items_data:

            qty = int(it.get('quantity', 1))
            price = float(it.get('unit_price', 0))
            name = (it.get('item_name') or '').strip()

            if not name or qty <= 0 or price < 0:
                return err('Invalid item data')

            total = round(qty * price, 2)

            subtotal += total

            parsed.append({
                'item_name': name,
                'quantity': qty,
                'unit_price': price,
                'total': total
            })

        subtotal = round(subtotal, 2)

        grand = round(max(subtotal - discount, 0), 2)

        pending = round(max(grand - paid, 0), 2)

        status = bill_status(grand, paid)

        cur = c.execute(
            """
            INSERT INTO bills
            (
                bill_number,
                bill_date,
                customer_id,
                subtotal,
                discount,
                grand_total,
                amount_paid,
                pending,
                status,
                notes
            )
            VALUES (?,?,?,?,?,?,?,?,?,?)
            """,
            (
                bill_no,
                bill_date,
                int(cid),
                subtotal,
                discount,
                grand,
                paid,
                pending,
                status,
                notes
            )
        )

        bid = cur.lastrowid

        for it in parsed:

            c.execute(
                """
                INSERT INTO bill_items
                (
                    bill_id,
                    item_name,
                    quantity,
                    unit_price,
                    total
                )
                VALUES (?,?,?,?,?)
                """,
                (
                    bid,
                    it['item_name'],
                    it['quantity'],
                    it['unit_price'],
                    it['total']
                )
            )

        if paid > 0:

            c.execute(
                """
                INSERT INTO payments
                (
                    bill_id,
                    amount,
                    method,
                    paid_at
                )
                VALUES (?,?,?,?)
                """,
                (
                    bid,
                    paid,
                    d.get('payment_method', 'cash'),
                    datetime.now().isoformat(timespec='seconds')
                )
            )

        c.commit()

        row = c.execute(
            """
            SELECT
                b.*,
                cu.name AS customer_name,
                cu.mobile AS customer_mobile
            FROM bills b
            LEFT JOIN customers cu
            ON b.customer_id = cu.id
            WHERE b.id=?
            """,
            (bid,)
        ).fetchone()

        return ok(dict(row), 201)

    finally:
        c.close()
    
@app.route('/update_bill/<int:i>',methods=['PUT'])
def update_bill(i):
    d=request.get_json(silent=True) or {}; items_data=d.get('items',[]); c=get_db()
    try:
        bill=c.execute('SELECT * FROM bills WHERE id=?',(i,)).fetchone()
        if not bill: return err('Not found',404)
        try: discount=float(d.get('discount',bill['discount'])); assert discount>=0
        except: return err('Invalid discount')
        notes=(d.get('notes') or '').strip(); subtotal=0.0; parsed=[]
        for it in items_data:
            qty=int(it.get('quantity',1)); price=float(it.get('unit_price',0)); name=(it.get('item_name') or '').strip()
            total=round(qty*price,2); subtotal+=total
            parsed.append({'item_name':name,'quantity':qty,'unit_price':price,'total':total})
        subtotal=round(subtotal,2); grand=round(max(subtotal-discount,0),2)
        paid_total=c.execute('SELECT COALESCE(SUM(amount),0) FROM payments WHERE bill_id=?',(i,)).fetchone()[0]
        pending=round(max(grand-paid_total,0),2); status=bill_status(grand,paid_total)
        c.execute("UPDATE bills SET subtotal=?,discount=?,grand_total=?,amount_paid=?,pending=?,status=?,notes=? WHERE id=?",(subtotal,discount,grand,paid_total,pending,status,notes,i))
        c.execute('DELETE FROM bill_items WHERE bill_id=?',(i,))
        for it in parsed:
            c.execute("INSERT INTO bill_items (bill_id,item_name,quantity,unit_price,total) VALUES (?,?,?,?,?)",(i,it['item_name'],it['quantity'],it['unit_price'],it['total']))
        c.commit(); return ok()
    finally: c.close()

@app.route('/delete_bill/<int:i>',methods=['DELETE'])
def delete_bill(i):
    c=get_db()
    try:
        if not c.execute('SELECT id FROM bills WHERE id=?',(i,)).fetchone(): return err('Not found',404)
        c.execute('DELETE FROM bills WHERE id=?',(i,)); c.commit(); return ok()
    finally: c.close()

# ══════════════════════════════════════════════════════════════════════════════
# PAYMENTS
# ══════════════════════════════════════════════════════════════════════════════

@app.route('/add_payment',methods=['POST'])
def add_payment():
    d=request.get_json(silent=True) or {}; bill_id=d.get('bill_id')
    if not bill_id: return err('bill_id required')
    try: amount=float(d.get('amount')); assert amount>0
    except: return err('amount must be positive')
    method=(d.get('method') or 'cash'); note=(d.get('note') or '').strip(); c=get_db()
    try:
        bill=c.execute('SELECT * FROM bills WHERE id=?',(bill_id,)).fetchone()
        if not bill: return err('Bill not found',404)
        c.execute("INSERT INTO payments (bill_id,amount,method,note,paid_at) VALUES (?,?,?,?,?)",(bill_id,amount,method,note,datetime.now().isoformat(timespec='seconds')))
        total_paid=c.execute('SELECT COALESCE(SUM(amount),0) FROM payments WHERE bill_id=?',(bill_id,)).fetchone()[0]
        pending=round(max(bill['grand_total']-total_paid,0),2); status=bill_status(bill['grand_total'],total_paid)
        c.execute('UPDATE bills SET amount_paid=?,pending=?,status=? WHERE id=?',(total_paid,pending,status,bill_id))
        c.commit(); return ok({'success':True,'total_paid':total_paid,'pending':pending,'status':status})
    finally: c.close()

@app.route('/pending_bills')
def get_pending_bills():
    q=request.args.get('q','').strip(); c=get_db()
    try:
        sql="SELECT b.*,cu.name AS customer_name,cu.mobile AS customer_mobile FROM bills b JOIN customers cu ON b.customer_id=cu.id WHERE b.status!='paid'"
        params=[]
        if q:
            sql+=" AND (cu.name LIKE ? OR cu.mobile LIKE ? OR b.bill_number LIKE ?)"
            like=f'%{q}%'; params=[like,like,like]
        sql+=" ORDER BY b.pending DESC"
        rows=c.execute(sql,params).fetchall(); return ok([dict(r) for r in rows])
    finally: c.close()

# ══════════════════════════════════════════════════════════════════════════════
# SETTINGS
# ══════════════════════════════════════════════════════════════════════════════

@app.route('/settings')
def get_settings():
    c=get_db(); rows=c.execute('SELECT key,value FROM settings').fetchall(); c.close()
    return ok({r['key']:r['value'] for r in rows})

@app.route('/settings',methods=['POST'])
def save_settings():
    d=request.get_json(silent=True) or {}; c=get_db()
    try:
        for key in ['shop_name','shop_address','shop_mobile','shop_email','shop_gst','shop_logo','bill_prefix']:
            if key in d:
                val=d[key]
                if key=='shop_logo' and val and val.startswith('data:'):
                    val=save_image(val,'shop_logo') or ''
                c.execute("INSERT OR REPLACE INTO settings (key,value) VALUES (?,?)",(key,val or ''))
        c.commit(); return ok()
    finally: c.close()

# ══════════════════════════════════════════════════════════════════════════════
# GLOBAL SEARCH
# ══════════════════════════════════════════════════════════════════════════════

@app.route('/search')
def global_search():
    q=request.args.get('q','').strip()
    if len(q)<2: return ok({'customers':[],'bills':[],'items':[]})
    like=f'%{q}%'; c=get_db()
    try:
        custs=c.execute("SELECT id,name,mobile,city FROM customers WHERE name LIKE ? OR mobile LIKE ? LIMIT 8",(like,like)).fetchall()
        bills=c.execute("SELECT b.id,b.bill_number,b.bill_date,b.grand_total,b.status,cu.name AS customer_name FROM bills b JOIN customers cu ON b.customer_id=cu.id WHERE b.bill_number LIKE ? OR cu.name LIKE ? OR cu.mobile LIKE ? LIMIT 8",(like,like,like)).fetchall()
        items=c.execute("SELECT id,name,selling_price,quantity FROM items WHERE name LIKE ? LIMIT 8",(like,)).fetchall()
        return ok({'customers':[dict(r) for r in custs],'bills':[dict(r) for r in bills],'items':[dict(r) for r in items]})
    finally: c.close()

# ══════════════════════════════════════════════════════════════════════════════
# WHATSAPP MESSAGE
# ══════════════════════════════════════════════════════════════════════════════
@app.route('/whatsapp_message/<int:bid>')
def whatsapp_message(bid):
    c = get_db()

    try:
        bill = c.execute("""
            SELECT b.*,
                   cu.name AS customer_name,
                   cu.mobile AS customer_mobile
            FROM bills b
            JOIN customers cu ON b.customer_id = cu.id
            WHERE b.id=?
        """, (bid,)).fetchone()

        if not bill:
            return err('Bill not found', 404)

        items = c.execute("""
            SELECT *
            FROM bill_items
            WHERE bill_id=?
            ORDER BY id
        """, (bid,)).fetchall()

        shop = {
            r['key']: r['value']
            for r in c.execute(
                'SELECT key,value FROM settings'
            ).fetchall()
        }

        sn = shop.get('shop_name', 'SpareTrack')

        items_text = ""

        for idx, item in enumerate(items, start=1):
            items_text += (
                f"{idx}. {item['item_name']}\n"
                f"   Qty : {item['quantity']}\n"
                f"   Rate : ₹{item['unit_price']:,.2f}\n"
                f"   Amount : ₹{item['total']:,.2f}\n\n"
            )

        status_text = {
            'paid': 'Fully Paid ✅',
            'partial': 'Partially Paid ⏳',
            'pending': 'Payment Pending ⚠️'
        }

        msg = (
            f"🏍️*{sn}*\n"
            f"━━━━━━━━━━━━━━━━━━\n"
            f"Customer : *{bill['customer_name']}*\n"
            f"Bill No : {bill['bill_number']}\n"
            f"Date : {bill['bill_date']}\n\n"

            f"🛒 *Items Purchased*\n"
            f"{items_text}"

            f"━━━━━━━━━━━━━━━━━━\n"
            f"💰 *Payment Summary*\n"
            f"Subtotal : ₹{bill['subtotal']:,.2f}\n"
            f"Discount : ₹{bill['discount']:,.2f}\n"
            f"Total : ₹{bill['grand_total']:,.2f}\n"
            f"Paid : ₹{bill['amount_paid']:,.2f}\n"
            f"Pending : ₹{bill['pending']:,.2f}\n"
            f"Status : {status_text.get(bill['status'])}\n"
            f"━━━━━━━━━━━━━━━━━━\n"
            f"Thank you "
        )

        mob = (bill['customer_mobile'] or '').replace(' ', '').replace('-', '')

        if mob.startswith('0'):
            mob = '91' + mob[1:]
        elif not mob.startswith('91'):
            mob = '91' + mob

        from urllib.parse import quote

        wa_url = f"https://wa.me/{mob}?text={quote(msg)}"

        return ok({
            'message': msg,
            'whatsapp_url': wa_url,
            'mobile': mob
        })

    finally:
        c.close()
# ══════════════════════════════════════════════════════════════════════════════
# PDF INVOICE
# ══════════════════════════════════════════════════════════════════════════════

@app.route('/invoice_pdf/<int:bid>')
def invoice_pdf(bid):
    if not HAS_REPORTLAB: return err('ReportLab not installed. Run: pip install reportlab',500)
    c=get_db()
    try:
        bill=c.execute("SELECT b.*,cu.name AS customer_name,cu.mobile AS customer_mobile,cu.email AS customer_email,cu.address AS customer_address,cu.city AS customer_city FROM bills b JOIN customers cu ON b.customer_id=cu.id WHERE b.id=?",(bid,)).fetchone()
        if not bill: return err('Bill not found',404)
        items=c.execute('SELECT * FROM bill_items WHERE bill_id=? ORDER BY id',(bid,)).fetchall()
        payments=c.execute('SELECT * FROM payments WHERE bill_id=? ORDER BY paid_at',(bid,)).fetchall()
        shop={r['key']:r['value'] for r in c.execute('SELECT key,value FROM settings').fetchall()}
    finally: c.close()

    buf=io.BytesIO()
    doc=SimpleDocTemplate(buf,pagesize=A4,rightMargin=15*mm,leftMargin=15*mm,topMargin=15*mm,bottomMargin=15*mm)
    story=[]
    sn_style=ParagraphStyle('SN',fontSize=15,leading=18,fontName='Helvetica-Bold',textColor=colors.HexColor('#1e40af'))
    si_style=ParagraphStyle('SI',fontSize=8,leading=10)
    inv_style=ParagraphStyle('IN',fontSize=18,fontName='Helvetica-Bold',textColor=colors.HexColor('#6366f1'),alignment=TA_RIGHT)
    id_style=ParagraphStyle('ID',fontSize=9,alignment=TA_RIGHT,leading=13)
    sc_hex={'paid':'#16a34a','partial':'#d97706','pending':'#dc2626'}.get(bill['status'],'#374151')
    left=[Paragraph(shop.get('shop_name','SpareTrack'),sn_style)]
    for k,pfx in [('shop_address',''),('shop_mobile','📞 '),('shop_email','✉ '),('shop_gst','GST: ')]:
        if shop.get(k): left.append(Paragraph(pfx+shop[k],si_style))
    right = [
    Paragraph('INVOICE', inv_style),

    Paragraph(
        f"Bill No: <b>{bill['bill_number']}</b>",
        ParagraphStyle(
            'BN',
            fontSize=10,
            alignment=TA_RIGHT,
            textColor=colors.HexColor('#4f46e5')
        )
    ),

    Paragraph(
        f"Date: {bill['bill_date']}",
        id_style
    ),

    Spacer(1,4),

    Paragraph(
        f"<font color='{sc_hex}'><b>{bill['status'].upper()}</b></font>",
        ParagraphStyle(
            'ST',
            fontSize=10,
            alignment=TA_RIGHT
        )
    )
]
     
    ht = Table(
    [[left, right]],
    colWidths=[120*mm, 60*mm]
   )
   
    
    ht = Table(
    [[left, right]],
    colWidths=[120*mm, 60*mm]
    )

    ht.setStyle(
    TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'TOP')
    ])
)

    story.append(ht)
    story.append(HRFlowable(width='100%',thickness=2,color=colors.HexColor('#4f46e5'),spaceAfter=8))
    cs=ParagraphStyle('CS',fontSize=9,leading=13); ch=ParagraphStyle('CH',fontSize=10,fontName='Helvetica-Bold',textColor=colors.HexColor('#374151'))
    crow=[[Paragraph('<b>BILL TO</b>',ch),'',Paragraph('<b>BILL INFO</b>',ch),''],[Paragraph('Customer:',cs),Paragraph(bill['customer_name'],cs),Paragraph('Bill No:',cs),Paragraph(bill['bill_number'],cs)],[Paragraph('Mobile:',cs),Paragraph(bill['customer_mobile'] or '—',cs),Paragraph('Date:',cs),Paragraph(bill['bill_date'],cs)]]
    if bill['customer_address'] or bill['customer_city']:
        addr=' '.join(filter(None,[bill['customer_address'],bill['customer_city']]))
        crow.append([Paragraph('Address:',cs),Paragraph(addr,cs),'',''])
    ct=Table(crow,colWidths=[24*mm,72*mm,24*mm,60*mm])
    ct.setStyle(TableStyle([('BACKGROUND',(0,0),(0,0),colors.HexColor('#eff6ff')),('BACKGROUND',(2,0),(2,0),colors.HexColor('#eff6ff')),('SPAN',(0,0),(1,0)),('SPAN',(2,0),(3,0)),('BOTTOMPADDING',(0,0),(-1,-1),4),('TOPPADDING',(0,0),(-1,-1),4),('GRID',(0,0),(-1,-1),0.3,colors.HexColor('#e5e7eb'))]))
    story.append(ct); story.append(Spacer(1,8))
    hs=ParagraphStyle('HS',fontSize=9,fontName='Helvetica-Bold',textColor=colors.white)
    is2=ParagraphStyle('IS',fontSize=9); ns=ParagraphStyle('NS',fontSize=9,alignment=TA_RIGHT)
    td=[[Paragraph(h,hs) for h in ['#','Item Description','Qty','Unit Price','Total']]]
    for idx,it in enumerate(items,1):
        td.append([Paragraph(str(idx),is2),Paragraph(it['item_name'],is2),Paragraph(str(it['quantity']),ns),Paragraph(f"₹{it['unit_price']:,.2f}",ns),Paragraph(f"₹{it['total']:,.2f}",ns)])
    it_tbl=Table(td,colWidths=[10*mm,80*mm,18*mm,30*mm,30*mm])
    it_tbl.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#4f46e5')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f8faff')]),('GRID',(0,0),(-1,-1),0.3,colors.HexColor('#e5e7eb')),('BOTTOMPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),5),('ALIGN',(2,0),(-1,-1),'RIGHT')]))
    story.append(it_tbl); story.append(Spacer(1,6))
    ts=ParagraphStyle('TS',fontSize=9); tbs=ParagraphStyle('TBS',fontSize=9,fontName='Helvetica-Bold')
    tots=[[Paragraph('Subtotal:',ts),Paragraph(f"₹{bill['subtotal']:,.2f}",ns)],[Paragraph('Discount:',ts),Paragraph(f"₹{bill['discount']:,.2f}",ns)],[Paragraph('<b>Grand Total:</b>',ParagraphStyle('GT',fontSize=11,fontName='Helvetica-Bold')),Paragraph(f"<b>₹{bill['grand_total']:,.2f}</b>",ParagraphStyle('GTV',fontSize=11,fontName='Helvetica-Bold',alignment=TA_RIGHT,textColor=colors.HexColor('#4f46e5')))],[Paragraph('Amount Paid:',ts),Paragraph(f"₹{bill['amount_paid']:,.2f}",ns)],[Paragraph('<b>Pending:</b>',tbs),Paragraph(f"<b>₹{bill['pending']:,.2f}</b>",ParagraphStyle('PV',fontSize=9,fontName='Helvetica-Bold',alignment=TA_RIGHT,textColor=colors.HexColor(sc_hex)))]]
    story.append(Table([[Spacer(1,1),Table(tots,colWidths=[45*mm,35*mm])]],colWidths=[100*mm,80*mm]))
    if payments:
        story.append(Spacer(1,8)); story.append(HRFlowable(width='100%',thickness=0.5,color=colors.HexColor('#e5e7eb')))
        story.append(Paragraph('<b>Payment History</b>',ParagraphStyle('PH',fontSize=10,fontName='Helvetica-Bold',spaceBefore=6)))
        ph=[[Paragraph(h,hs) for h in ['Date','Method','Amount','Note']]]
        for p in payments:
            ph.append([Paragraph(p['paid_at'][:10],is2),Paragraph(p['method'].title(),is2),Paragraph(f"₹{p['amount']:,.2f}",ns),Paragraph(p['note'] or '—',is2)])
        pt=Table(ph,colWidths=[30*mm,30*mm,30*mm,90*mm])
        pt.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#6366f1')),('GRID',(0,0),(-1,-1),0.3,colors.HexColor('#e5e7eb')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f8faff')]),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]))
        story.append(pt)
    story.append(Spacer(1,12)); story.append(HRFlowable(width='100%',thickness=0.5,color=colors.HexColor('#e5e7eb')))
    story.append(Paragraph(f"Thank you for your business! — {shop.get('shop_name','')}",ParagraphStyle('F',fontSize=9,textColor=colors.HexColor('#6b7280'),alignment=TA_CENTER,spaceBefore=4)))
    doc.build(story); buf.seek(0)
    return Response(buf.read(),mimetype='application/pdf',headers={'Content-Disposition':f'attachment; filename="invoice_{bill["bill_number"]}.pdf"'})

# ══════════════════════════════════════════════════════════════════════════════
# REPORTS
# ══════════════════════════════════════════════════════════════════════════════

def _rdata(rt,start=None,end=None):
    today=date.today().isoformat()
    if not end: end=today
    if not start:
        if rt=='daily': start=today
        elif rt=='weekly': start=(date.today()-timedelta(days=6)).isoformat()
        elif rt=='monthly': start=date.today().replace(day=1).isoformat()
        else: start='2000-01-01'
    c=get_db()
    try:
        bills=c.execute("SELECT b.*,cu.name AS customer_name,cu.mobile AS customer_mobile FROM bills b JOIN customers cu ON b.customer_id=cu.id WHERE b.bill_date BETWEEN ? AND ? ORDER BY b.bill_date DESC",(start,end)).fetchall()
        br=[dict(b) for b in bills]
        return {'bills':br,'summary':{'start_date':start,'end_date':end,'bill_count':len(br),'total_revenue':round(sum(b['grand_total'] for b in br),2),'total_paid':round(sum(b['amount_paid'] for b in br),2),'total_pending':round(sum(b['pending'] for b in br),2),'report_type':rt}}
    finally: c.close()

@app.route('/report/data')
def report_data_api():
    return ok(_rdata(request.args.get('type','daily'),request.args.get('start') or None,request.args.get('end') or None))

@app.route('/report/pdf')
def report_pdf():
    if not HAS_REPORTLAB: return err('ReportLab not installed',500)
    rt=request.args.get('type','daily'); data=_rdata(rt,request.args.get('start') or None,request.args.get('end') or None)
    s=data['summary']; bills=data['bills']
    c=get_db(); shop={r['key']:r['value'] for r in c.execute('SELECT key,value FROM settings').fetchall()}; c.close()
    buf=io.BytesIO()
    doc=SimpleDocTemplate(buf,pagesize=A4,rightMargin=15*mm,leftMargin=15*mm,topMargin=15*mm,bottomMargin=15*mm)
    story=[Paragraph(shop.get('shop_name','SpareTrack'),ParagraphStyle('T',fontSize=18,fontName='Helvetica-Bold',textColor=colors.HexColor('#1e40af'),alignment=TA_CENTER)),
           Paragraph(f"{rt.title()} Sales Report: {s['start_date']} → {s['end_date']}",ParagraphStyle('S',fontSize=12,alignment=TA_CENTER,spaceBefore=4)),
           HRFlowable(width='100%',thickness=2,color=colors.HexColor('#4f46e5'),spaceAfter=10)]
    hs=ParagraphStyle('HS',fontSize=9,fontName='Helvetica-Bold',textColor=colors.white)
    is2=ParagraphStyle('IS',fontSize=8); rs=ParagraphStyle('RS',fontSize=8,alignment=TA_RIGHT)
    sd=[[Paragraph(h,ParagraphStyle('SC',fontSize=10,alignment=TA_CENTER)) for h in ['Bills','Revenue','Collected','Pending']],
        [Paragraph(str(s['bill_count']),ParagraphStyle('SV',fontSize=18,fontName='Helvetica-Bold',alignment=TA_CENTER,textColor=colors.HexColor('#4f46e5'))),
         Paragraph(f"₹{s['total_revenue']:,.0f}",ParagraphStyle('SV2',fontSize=14,fontName='Helvetica-Bold',alignment=TA_CENTER,textColor=colors.HexColor('#16a34a'))),
         Paragraph(f"₹{s['total_paid']:,.0f}",ParagraphStyle('SV3',fontSize=14,fontName='Helvetica-Bold',alignment=TA_CENTER,textColor=colors.HexColor('#0284c7'))),
         Paragraph(f"₹{s['total_pending']:,.0f}",ParagraphStyle('SV4',fontSize=14,fontName='Helvetica-Bold',alignment=TA_CENTER,textColor=colors.HexColor('#dc2626')))]]
    st=Table(sd,colWidths=[44*mm]*4)
    st.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eff6ff')),('BOX',(0,0),(-1,-1),1,colors.HexColor('#bfdbfe')),('INNERGRID',(0,0),(-1,-1),0.5,colors.HexColor('#dbeafe')),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)]))
    story.append(st); story.append(Spacer(1,10))
    td=[[Paragraph(h,hs) for h in ['Bill No','Date','Customer','Total','Paid','Pending','Status']]]
    for b in bills:
        sc2={'paid':'#16a34a','partial':'#d97706','pending':'#dc2626'}.get(b['status'],'#374151')
        td.append([Paragraph(b['bill_number'],is2),Paragraph(b['bill_date'],is2),Paragraph(b['customer_name'][:20],is2),Paragraph(f"₹{b['grand_total']:,.0f}",rs),Paragraph(f"₹{b['amount_paid']:,.0f}",rs),Paragraph(f"₹{b['pending']:,.0f}",rs),Paragraph(f"<font color='{sc2}'>{b['status'].title()}</font>",is2)])
    tbl=Table(td,colWidths=[28*mm,22*mm,40*mm,24*mm,24*mm,24*mm,18*mm])
    tbl.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#4f46e5')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f8faff')]),('GRID',(0,0),(-1,-1),0.3,colors.HexColor('#e5e7eb')),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]))
    story.append(tbl); doc.build(story); buf.seek(0)
    return Response(buf.read(),mimetype='application/pdf',headers={'Content-Disposition':f'attachment; filename="report_{rt}_{s["start_date"]}.pdf"'})

@app.route('/report/excel')
def report_excel():
    if not HAS_OPENPYXL: return err('openpyxl not installed',500)
    rt=request.args.get('type','daily'); data=_rdata(rt,request.args.get('start') or None,request.args.get('end') or None)
    s=data['summary']; bills=data['bills']
    wb=openpyxl.Workbook(); ws=wb.active; ws.title=f"{rt.title()} Report"
    hf=PatternFill("solid",fgColor="4F46E5"); hfnt=Font(bold=True,color="FFFFFF",size=11)
    bold=Font(bold=True); center=Alignment(horizontal='center')
    thin=Side(border_style='thin',color='E5E7EB'); bdr=Border(left=thin,right=thin,top=thin,bottom=thin)
    ws.merge_cells('A1:H1'); ws['A1']=f"{rt.title()} Sales Report — {s['start_date']} to {s['end_date']}"; ws['A1'].font=Font(bold=True,size=14,color="1E40AF"); ws['A1'].alignment=center
    ws.append([]); ws.append(['Summary','','','','','','',''])
    for k,v in [['Bills',s['bill_count']],['Total Revenue',f"₹{s['total_revenue']:,.2f}"],['Total Paid',f"₹{s['total_paid']:,.2f}"],['Total Pending',f"₹{s['total_pending']:,.2f}"]]:
        ws.append([k,v]); ws.cell(ws.max_row,1).font=bold
    ws.append([])
    hdrs=['Bill Number','Date','Customer','Mobile','Total','Paid','Pending','Status']
    ws.append(hdrs)
    for col,h in enumerate(hdrs,1):
        cell=ws.cell(ws.max_row,col); cell.fill=hf; cell.font=hfnt; cell.alignment=center; cell.border=bdr
    sc_map={'paid':'D1FAE5','partial':'FEF3C7','pending':'FEE2E2'}
    for b in bills:
        ws.append([b['bill_number'],b['bill_date'],b['customer_name'],b['customer_mobile'],b['grand_total'],b['amount_paid'],b['pending'],b['status'].title()])
        fc=sc_map.get(b['status'],'FFFFFF')
        for col in range(1,9):
            ws.cell(ws.max_row,col).fill=PatternFill("solid",fgColor=fc); ws.cell(ws.max_row,col).border=bdr
    for col,w in zip('ABCDEFGH',[16,12,24,14,12,12,12,10]): ws.column_dimensions[col].width=w
    buf=io.BytesIO(); wb.save(buf); buf.seek(0)
    return Response(buf.read(),mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="report_{rt}_{s["start_date"]}.xlsx"'})

# ══════════════════════════════════════════════════════════════════════════════
# BACKUP & RESTORE
# ══════════════════════════════════════════════════════════════════════════════

@app.route('/backup')
def backup_db():
    if not os.path.exists(DB_PATH): return err('No database found',404)
    return send_file(DB_PATH,as_attachment=True,download_name=f"sparetrack_backup_{date.today().isoformat()}.db",mimetype='application/octet-stream')

@app.route('/restore',methods=['POST'])
def restore_db():
    if 'file' not in request.files: return err('No file uploaded')
    f=request.files['file']
    if not f.filename.endswith('.db'): return err('File must be .db')
    if os.path.exists(DB_PATH): shutil.copy2(DB_PATH,os.path.join(BACKUP_DIR,f"pre_restore_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"))
    f.save(DB_PATH); return ok({'message':'Database restored successfully'})

# ─── Entry ────────────────────────────────────────────────────────────────────

def run_server(data_dir=None, port=5000):
    """
    Entry point used on Android (called from MainActivity via Chaquopy).
    `data_dir` is the app's private, writable storage folder on the phone
    (Context.getFilesDir()) — this is where database.db, images and backups
    are kept, so all data survives app restarts and works fully offline
    (127.0.0.1 loopback needs no internet connection).
    Desktop/Electron continue to use the `if __name__=='__main__'` path below
    and are unaffected by this function.
    """
    global DATA_DIR, DB_PATH, IMAGES_DIR, BACKUP_DIR
    if data_dir:
        DATA_DIR   = data_dir
        DB_PATH    = os.path.join(DATA_DIR, 'database.db')
        IMAGES_DIR = os.path.join(DATA_DIR, 'static', 'images')
        BACKUP_DIR = os.path.join(DATA_DIR, 'backups')
        os.makedirs(IMAGES_DIR, exist_ok=True)
        os.makedirs(BACKUP_DIR, exist_ok=True)
    init_db()
    app.run(host='127.0.0.1', port=port, debug=False, use_reloader=False, threaded=True)

if __name__ == '__main__':
    init_db()
    port=int(os.environ.get('FLASK_PORT',5000))
    print(f"\n🚀  SpareTrack  →  http://localhost:{port}/")
    print(f"    PDF: {'✅' if HAS_REPORTLAB else '❌ pip install reportlab'}")
    print(f"    Excel: {'✅' if HAS_OPENPYXL else '❌ pip install openpyxl'}")
    app.run(host='0.0.0.0',port=port,debug=False,use_reloader=False)
